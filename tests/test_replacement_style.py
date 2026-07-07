import time
from io import BytesIO

import pytest
from doctotext import DOCX_MIME, load_document
from docx import Document
from fastapi.testclient import TestClient

from anonimizator3000.config import (
    DEFAULT_REPLACEMENT_STYLE,
    normalize_replacement_style,
    settings_from_env,
)
from anonimizator3000.main import app

_SAMPLE = "Jan Kowalski ma PESEL 44051401359."


def _docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_normalize_replacement_style_accepts_known_values() -> None:
    assert normalize_replacement_style("labels") == "labels"
    assert normalize_replacement_style("MASK") == "mask"
    assert normalize_replacement_style("  Labels  ") == "labels"


def test_normalize_replacement_style_falls_back_to_default() -> None:
    assert normalize_replacement_style(None) == DEFAULT_REPLACEMENT_STYLE
    assert normalize_replacement_style("nonsense") == DEFAULT_REPLACEMENT_STYLE
    assert normalize_replacement_style("nonsense", "labels") == "labels"


def test_settings_from_env_reads_replacement_style(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("ANON_REPLACEMENT_STYLE", "labels")
    assert settings_from_env().replacement_style == "labels"

    monkeypatch.setenv("ANON_REPLACEMENT_STYLE", "bogus")
    assert settings_from_env().replacement_style == DEFAULT_REPLACEMENT_STYLE


def test_index_hides_style_choice() -> None:
    with TestClient(app) as client:
        response = client.get("/")

    assert response.status_code == 200
    assert 'name="style"' not in response.text
    assert "Styl anonimizacji" not in response.text


def test_labels_style_produces_category_placeholders() -> None:
    text = _run_upload("labels")

    assert "Jan Kowalski" not in text
    assert "44051401359" not in text
    assert "[OSOBA_1]" in text
    assert "[PESEL_1]" in text
    assert "****" not in text


def test_default_style_produces_fixed_mask() -> None:
    text = _run_upload(None)

    assert "Jan Kowalski" not in text
    assert "44051401359" not in text
    assert "****" in text
    assert "[OSOBA_1]" not in text


def _run_upload(style: str | None) -> str:
    files = {"document": ("sample.docx", _docx_bytes(_SAMPLE), DOCX_MIME)}
    data = {"style": style} if style is not None else None
    with TestClient(app) as client:
        response = client.post("/jobs", files=files, data=data)
        assert response.status_code == 200
        job_id = response.text.split("/jobs/", 1)[1].split('"', 1)[0]

        for _ in range(50):
            status = client.get(f"/jobs/{job_id}")
            assert status.status_code == 200
            if "Gotowe" in status.text:
                break
            time.sleep(0.05)
        else:
            raise AssertionError("Job did not finish")

        download = client.get(f"/jobs/{job_id}/download")
        assert download.status_code == 200
        document = load_document("wynik.docx", DOCX_MIME, download.content)
        return "\n".join(document.texts)
